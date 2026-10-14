import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn
import os

def create_tseda_docx():
    doc = docx.Document()

    # Set 0.75 in margins
    sections = doc.sections
    for s in sections:
        s.top_margin = Inches(0.75)
        s.bottom_margin = Inches(0.75)
        s.left_margin = Inches(0.75)
        s.right_margin = Inches(0.75)

    def set_cell_background(cell, fill_hex):
        shading_elm = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
        cell._tc.get_or_add_tcPr().append(shading_elm)

    def add_heading_title(text, level=1):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        run.bold = True
        run.font.name = 'Arial'
        if level == 1:
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0x4A, 0x15, 0x19)
        elif level == 2:
            run.font.size = Pt(12)
            run.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.space_before = Pt(6)
        return p

    def add_section_header(title):
        p = doc.add_paragraph()
        run = p.add_run(title)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(11)
        run.font.color.rgb = RGBColor(0x4A, 0x15, 0x19)
        p.paragraph_format.space_before = Pt(12)
        p.paragraph_format.space_after = Pt(4)
        return p

    def add_sub_header(title):
        p = doc.add_paragraph()
        run = p.add_run(title)
        run.bold = True
        run.font.name = 'Arial'
        run.font.size = Pt(10)
        run.font.color.rgb = RGBColor(0x37, 0x41, 0x51)
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        return p

    def format_table(table, col_widths=None):
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, row in enumerate(table.rows):
            # CantSplit
            trPr = row._tr.get_or_add_trPr()
            trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
            
            for j, cell in enumerate(row.cells):
                if col_widths and j < len(col_widths):
                    cell.width = Inches(col_widths[j])
                cell.paragraphs[0].paragraph_format.space_before = Pt(2)
                cell.paragraphs[0].paragraph_format.space_after = Pt(2)
                if i == 0:
                    set_cell_background(cell, "F3F4F6")
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.bold = True
                            run.font.size = Pt(9.5)
                            run.font.name = 'Arial'
                else:
                    for p in cell.paragraphs:
                        for run in p.runs:
                            run.font.size = Pt(9)
                            run.font.name = 'Arial'

    # Title
    p1 = doc.add_paragraph()
    p1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p1.add_run("Thiagarajar College of Engineering: Madurai 625 015")
    r1.bold = True
    r1.font.size = Pt(13)
    r1.font.name = 'Arial'
    r1.font.color.rgb = RGBColor(0x4A, 0x15, 0x19)

    p2 = doc.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = p2.add_run("Template of Metrics for Appraisal of Performance of\nAssistant Professor, Associate Professors and Professors\nTSEDA (Department of Architecture)")
    r2.bold = True
    r2.font.size = Pt(11)
    r2.font.name = 'Arial'
    p2.paragraph_format.space_after = Pt(10)

    # Faculty Info Table
    info_table = doc.add_table(rows=7, cols=2)
    info_rows = [
        ("Name of the Faculty", ""),
        ("Designation", ""),
        ("Department", "Architecture (TSEDA)"),
        ("Highest Qualification", ""),
        ("Date of Joining at TCE", ""),
        ("Date of Joining in the Present Post", ""),
        ("Number of Years of Experience at TCE", "")
    ]
    for idx, (label, val) in enumerate(info_rows):
        row = info_table.rows[idx]
        row.cells[0].paragraphs[0].add_run(label).bold = True
        row.cells[1].paragraphs[0].add_run(val)
    format_table(info_table, [3.0, 4.0])

    # SECTION I: TEACHING AND LEARNING
    add_section_header("SECTION I: TEACHING AND LEARNING")

    add_sub_header("1.1 Courses Handled (Theory / Lab / Integrated)")
    t = doc.add_table(rows=3, cols=4)
    headers = ["Course Code", "Course Name", "Type (Theory/Lab/Integrated)", "Semester"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [1.5, 3.0, 1.5, 1.0])

    add_sub_header("1.2 Maintenance of Course File")
    t = doc.add_table(rows=3, cols=4)
    headers = ["Course Code", "Course Name", "Semester", "Course File Completed? (Yes/No)"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [1.5, 3.0, 1.0, 1.5])

    add_sub_header("1.3 Course Designed / Revised")
    t = doc.add_table(rows=2, cols=3)
    headers = ["Course Code", "Course Name", "Remarks"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [1.5, 3.0, 2.5])

    add_sub_header("1.4 Value-Added / Industry supported / Skill Courses / Documentation and Skill Workshops / Design-Build Workshops Organized")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Course Name", "Particulars", "Semester", "No. of Students"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.5, 2.5, 1.0, 1.0])

    add_sub_header("1.5.1 Innovative Teaching Methods adopted - Live Design Studio Pedagogy / Community-Based Learning / Participatory Design Exercises / Design Thinking Workshops / Urban Field Documentation / Climate Simulation Integration / Computational Design Integration / Design-Build Exercises / Interdisciplinary Studios / Joint Studios / Joint Teaching with HLI")
    t = doc.add_table(rows=2, cols=3)
    headers = ["Course Code", "Innovative Teaching Method", "Evidence / Particulars"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [1.5, 3.0, 2.5])

    add_sub_header("1.5.2 Studio-Based Teaching & Design Education - Achievement of Studio Focus Area / Studio Documentation & Process Recording / Open Reviews / Participation in Courses / Student Work Exhibitions / Beyond-Syllabus Exploration / Integration of SDGs / Climate / Community / Computational Tools")
    t = doc.add_table(rows=2, cols=3)
    headers = ["Course Name", "Activity", "Evidence"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.5, 2.5, 2.0])

    add_sub_header("1.6 Educational Tour And Case Study Visits")
    t = doc.add_table(rows=2, cols=5)
    headers = ["No. of Days", "Places of Visit", "Activity", "Evidence – Approval Letters / Geotagged Photos", "Number of Students & Year"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [1.0, 1.8, 1.5, 1.7, 1.0])

    add_sub_header("1.7 Mentoring System Effectiveness - Fast Learner / Slow Learner / Competition / Portfolio Development / Student Research / Student Publication Mentoring")
    t = doc.add_table(rows=2, cols=5)
    headers = ["No. / Name of Mentees", "Type of Mentoring", "Batch (e.g. 2024-29)", "Description of Effectiveness", "HoD Review (Excellent/Good/Avg/Poor)"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [1.5, 1.5, 1.2, 1.8, 1.0])

    add_sub_header("1.8 NPTEL / SWAYAM Certification (Count)")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Course Name", "Platform (NPTEL/SWAYAM)", "Certification Type (Elite / Silver / Gold / Completed)", "Month, Year"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.5, 1.5, 2.0, 1.0])

    add_sub_header("1.9 Student Feedback")
    t = doc.add_table(rows=3, cols=4)
    headers = ["Course Code", "Course Name", "Semester", "Score"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [1.5, 3.5, 1.0, 1.0])

    add_sub_header("1.10 Result Analysis")
    t = doc.add_table(rows=3, cols=4)
    headers = ["Course Code", "Course Name", "Semester", "Pass Percentage (%)"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [1.5, 3.5, 1.0, 1.0])

    add_sub_header("1.11 CO Attainment Level (Average % Achieved)")
    t = doc.add_table(rows=3, cols=4)
    headers = ["Course Code", "Course Name", "Semester", "Avg. CO Attainment (%)"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [1.5, 3.5, 1.0, 1.0])

    # SECTION II: RESEARCH AND BOOK PUBLICATIONS
    add_section_header("SECTION II: RESEARCH AND BOOK PUBLICATIONS")

    add_sub_header("2.1 Research Paper — SCI/SCI-E/SSCI/Scopus indexed Q1/Q2/Q3 Journals")
    t = doc.add_table(rows=2, cols=7)
    headers = ["Name of Journal", "Title of Paper", "Quartile", "Vol & Issue", "Year", "Authors as in Paper", "Index"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [1.5, 1.8, 0.7, 0.8, 0.6, 1.0, 0.6])

    add_sub_header("2.2 Total Scopus Citations Received with TCE affiliations in the last three years")
    t = doc.add_table(rows=2, cols=2)
    headers = ["Calendar Years", "Total Citations Received"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [3.5, 3.5])

    add_sub_header("2.3 Total Q1 journal Citations (Scopus) Received with TCE affiliations in the last three years")
    t = doc.add_table(rows=2, cols=2)
    headers = ["Calendar Years", "Total Q1 journal Citations"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [3.5, 3.5])

    add_sub_header("2.4 Book / Book Chapter Publication")
    t = doc.add_table(rows=2, cols=6)
    headers = ["Sl.No.", "Name of Publisher", "Title", "Type (Book/Chapter)", "Month & Year", "Authors"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [0.6, 1.8, 2.0, 1.0, 0.8, 0.8])

    add_sub_header("2.5 International / National Conference Publication")
    t = doc.add_table(rows=2, cols=6)
    headers = ["Sl.No.", "Name of Conference", "Title of Paper", "Indexed (Scopus)", "Month & Year", "Authors"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [0.6, 1.8, 2.0, 1.0, 0.8, 0.8])

    add_sub_header("2.6 Creative Scholarship & Professional Writing - Design Magazine Articles / Architectural Critiques / Urban Commentary / Visual Essays / Heritage Documentation Publications / Professional Design Publications / Curated Publications")
    t = doc.add_table(rows=2, cols=6)
    headers = ["Sl.No.", "Name of Journal/Magazine", "Title of Paper", "Vol & Issue / Proof", "Year", "Authors"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [0.6, 1.8, 2.0, 1.0, 0.8, 0.8])

    add_sub_header("2.7 No. of PhD Scholars Pursuing")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Name of Research Scholar", "Date of Admission", "Full Time / Part Time", "Status (Ongoing / Submitted)"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.5, 1.5, 1.5, 1.5])

    add_sub_header("2.8 No. of PhD Scholars Awarded in the Year")
    t = doc.add_table(rows=2, cols=5)
    headers = ["Name of Research Scholar", "Year of Registration", "Full Time / Part Time", "Viva Voce Date", "University"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.0, 1.2, 1.2, 1.2, 1.4])

    # SECTION III: PATENTS AND INNOVATION
    add_section_header("SECTION III: PATENTS AND INNOVATION")

    add_sub_header("3.1 Number of Patents Published")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Patent Ref Number", "Title", "Name of Inventors", "Date Published"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [1.8, 2.4, 1.6, 1.2])

    add_sub_header("3.2 Number of Patents Granted")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Patent Ref Number", "Title", "Name of Inventors", "Date Granted"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [1.8, 2.4, 1.6, 1.2])

    add_sub_header("3.3 Number of Transfer of Technology")
    t = doc.add_table(rows=2, cols=3)
    headers = ["Title", "Industry Partner", "Amount (Rs.)"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [3.0, 2.5, 1.5])

    add_sub_header("3.4 Number of Prototype / Product Developed (with Students)")
    t = doc.add_table(rows=2, cols=3)
    headers = ["Title of Product", "Students Involved", "Academic Year"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [3.0, 2.5, 1.5])

    add_sub_header("3.5 Hackathon Mentoring & Prizes (with Students)")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Hackathon / Event Name", "Students Mentored", "Prize / Achievement", "Academic Year"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.5, 1.8, 1.5, 1.2])

    add_sub_header("3.6 Design Patents")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Patent Ref Number", "Title", "Name of Inventors", "Date Published"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [1.8, 2.4, 1.6, 1.2])

    # SECTION IV: SPONSORED RESEARCH AND CONSULTANCY
    add_section_header("SECTION IV: SPONSORED RESEARCH AND CONSULTANCY")

    add_sub_header("4.1 Sponsored Research Project — PI / Co-PI")
    t = doc.add_table(rows=2, cols=6)
    headers = ["Project Name", "Funding Agency", "Period", "Sanctioned Amount (Rs.)", "Role (PI/Co-PI)", "Status (Ongoing/Completed)"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [1.8, 1.5, 1.0, 1.1, 0.8, 0.8])

    add_sub_header("4.2 Consultancy Projects")
    t = doc.add_table(rows=2, cols=5)
    headers = ["Title of Consultancy Project", "Funding Agency", "Period", "Amount Generated (Rs.)", "Names of Faculty Involved"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.0, 1.5, 1.0, 1.2, 1.3])

    # SECTION V: INTERNATIONAL ENGAGEMENT & RANKINGS CONTRIBUTION
    add_section_header("SECTION V: INTERNATIONAL ENGAGEMENT & RANKINGS CONTRIBUTION")

    add_sub_header("5.1 International Engagement / MoU")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Partner Institution", "Country", "Nature of Collaboration", "Status"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.5, 1.5, 2.0, 1.0])

    add_sub_header("5.2 Visiting / Adjunct Position Abroad")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Institution", "Country", "Duration (Days)", "Period"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.5, 1.5, 1.5, 1.5])

    add_sub_header("5.3 Foreign Faculty / Student Hosted or Engaged")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Name", "Institution / Country", "Engagement Type", "Period"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.0, 2.0, 1.5, 1.5])

    add_sub_header("5.4 QS / THE Reputation Survey Nominations Contribution")
    t = doc.add_table(rows=2, cols=3)
    headers = ["Survey Name", "Contribution Details", "Evidence Submitted (Yes/No)"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.5, 3.0, 1.5])

    add_sub_header("5.5 NIRF Survey Nomination")
    t = doc.add_table(rows=2, cols=2)
    headers = ["Nomination Details", "Evidence Submitted (Yes/No)"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [5.0, 2.0])

    add_sub_header("5.6 International Design Studio or other Collaborations")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Institution", "Country", "Nature of Collaboration", "Duration (Days) / Period"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.2, 1.4, 2.0, 1.4])

    # SECTION VI: FACULTY DEVELOPMENT & PROFESSIONAL ACTIVITIES
    add_section_header("SECTION VI: FACULTY DEVELOPMENT & PROFESSIONAL ACTIVITIES")

    add_sub_header("6.1 FDP / STTP Attended (5 days and above)")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Name of Program", "Organizer", "Duration", "Date (From - To)"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.5, 2.0, 1.0, 1.5])

    add_sub_header("6.2 Program Organized")
    t = doc.add_table(rows=2, cols=5)
    headers = ["Name of Program", "No. of Days", "Date (From - To)", "Coordinator / Co-coordinator", "No. of Participants"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.2, 0.9, 1.3, 1.4, 1.2])

    add_sub_header("6.3 Resource Person / Keynote Speaker")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Event Name", "International / National", "Topic", "Date"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.2, 1.5, 2.1, 1.2])

    add_sub_header("6.4 Professional Society Membership")
    t = doc.add_table(rows=2, cols=3)
    headers = ["Name of Society", "Membership Type", "Status (Active/Inactive)"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [3.0, 2.2, 1.8])

    add_sub_header("6.5 Designation in Professional Body / Editorial Board")
    t = doc.add_table(rows=2, cols=3)
    headers = ["Name of Body", "Position Held", "Period"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [3.0, 2.2, 1.8])

    add_sub_header("6.6 MOOCs Content Developed (TCE MOOC)")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Course Name", "Number of Weeks", "No. of Co-Faculty Members", "Number of Takers (Internal, External)"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.5, 1.2, 1.5, 1.8])

    # SECTION VII: INDUSTRY INTERACTION & INTERNSHIP
    add_section_header("SECTION VII: INDUSTRY INTERACTION & INTERNSHIP")

    add_sub_header("7.1 Partial delivery of regular courses by industry experts")
    t = doc.add_table(rows=2, cols=5)
    headers = ["Course Code & Name", "Online/Offline", "Name of Industry", "Expert Name & Designation", "Date"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.0, 1.0, 1.5, 1.5, 1.0])

    add_sub_header("7.2 Accompanying Industrial Visit")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Visit Details", "Industry", "No. of Students", "Date"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.5, 2.0, 1.2, 1.3])

    add_sub_header("7.3 Faculty Internship in Industries")
    t = doc.add_table(rows=2, cols=3)
    headers = ["Industry Name", "Duration (Days)", "Purpose"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.5, 1.5, 3.0])

    add_sub_header("7.4 Employer / Alumni Engagement Activity")
    t = doc.add_table(rows=2, cols=3)
    headers = ["Activity Name", "Alumni / Employer Involved", "Date"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [3.0, 2.5, 1.5])

    # SECTION VIII: STUDENT DEVELOPMENT ACTIVITIES
    add_section_header("SECTION VIII: STUDENT DEVELOPMENT ACTIVITIES")

    add_sub_header("8.1 UG/PG Student Project Publication- Journal/ Conference (Scopus indexed)")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Title", "Student Name(s)", "Journal / Conference details", "Date of Publication"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.5, 1.8, 1.8, 0.9])

    add_sub_header("8.2 Hackathon / Competition Mentoring - NASA/ZONASA/ INSDAG/IGBC/ Professional Competitions / Student Competitions / NASA Related Contributions / Urban Design Competitions / Heritage Competitions")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Hackathon / Event Name", "Students Mentored", "Outcome", "Duration/Date"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.5, 1.8, 1.5, 1.2])

    add_sub_header("8.3 Startup / Incubation Support / Tech-Club")
    t = doc.add_table(rows=2, cols=3)
    headers = ["Activity / Startup Name", "Role", "Duration"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [3.0, 2.0, 2.0])

    add_sub_header("8.4 Student Exhibitions")
    t = doc.add_table(rows=2, cols=4)
    headers = ["Event Name", "Students Mentored", "Outcome", "Duration/Date"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [2.5, 1.8, 1.5, 1.2])

    # SECTION IX: INSTITUTIONAL CONTRIBUTION
    add_section_header("SECTION IX: INSTITUTIONAL CONTRIBUTION")

    add_sub_header("9.1 Department Level Activities")
    t = doc.add_table(rows=2, cols=3)
    headers = ["Activity Description", "Role (Major/Supporting)", "HoD Approval"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [4.0, 1.8, 1.2])

    add_sub_header("9.2 College Level Activities")
    t = doc.add_table(rows=2, cols=3)
    headers = ["Activity Description", "Role (Major/Supporting)", "Section Head Approval"]
    for j, h in enumerate(headers):
        t.rows[0].cells[j].paragraphs[0].add_run(h).bold = True
    format_table(t, [4.0, 1.8, 1.2])

    # Declaration & Signatures
    add_section_header("DECLARATION")
    p_dec = doc.add_paragraph()
    p_dec.add_run("I hereby declare that all the information furnished is true to the best of my knowledge and belief.")
    p_dec.paragraph_format.space_before = Pt(6)
    p_dec.paragraph_format.space_after = Pt(20)

    t_sig = doc.add_table(rows=1, cols=2)
    t_sig.rows[0].cells[0].paragraphs[0].add_run("Date: _____________\n\nPlace: Madurai").bold = True
    p_sig = t_sig.rows[0].cells[1].paragraphs[0]
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig.add_run("\n\nSignature of the Faculty\n\nEndorsement of HoD (TSEDA)").bold = True
    format_table(t_sig, [3.5, 3.5])

    out_path = os.path.join("public", "TSEDA Faculty Appraisal Template Aug 2026.docx")
    doc.save(out_path)
    print(f"✓ Saved TSEDA Word template to {out_path}")

    # Also save with standard lowercase alias
    alias_path = os.path.join("public", "tseda_faculty_appraisal_template.docx")
    doc.save(alias_path)
    print(f"✓ Saved alias to {alias_path}")

create_tseda_docx()
